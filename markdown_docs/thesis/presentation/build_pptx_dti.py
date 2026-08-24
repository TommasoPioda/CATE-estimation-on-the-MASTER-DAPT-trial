#!/usr/bin/env python3
"""Rebuild the CATE/MASTER DAPT defense deck on top of the official SUPSI-DTI
PowerPoint template (title-slide branding, footer/date/page-number placeholders,
theme colors and fonts), instead of a fully custom from-scratch deck."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image as PILImage

IMG = "/home/tpioda/Bachelor-Thesis/markdown_docs/thesis/images/"
TEMPLATE = "/home/tpioda/Bachelor-Thesis/markdown_docs/thesis/DTI PowerPoint Template.pptx"
OUT = "/home/tpioda/Bachelor-Thesis/markdown_docs/thesis/presentation/CATE_MASTER_DAPT_DTI.pptx"

# ------------------------------------------------------- template palette --
# Pulled from the template's theme (SUPSI): dk1/lt1/accent1/accent2/accent5,
# not invented. BLEED/ISCH are kept exactly as in the referenced chart PNGs
# so text call-outs stay color-consistent with the embedded figures.
BG          = RGBColor(0xFF, 0xFF, 0xFF)   # theme bg1
SURFACE     = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE2    = RGBColor(0xF2, 0xF3, 0xF4)
INK         = RGBColor(0x00, 0x00, 0x00)   # theme dk1
INK_SOFT    = RGBColor(0x3C, 0x3C, 0x3C)   # theme dk2
INK_FAINT   = RGBColor(0x7A, 0x7A, 0x7A)
LINE        = RGBColor(0xD2, 0xD2, 0xD2)   # theme lt2
TITLE_BLUE  = RGBColor(0x28, 0x38, 0xC8)   # theme accent2 (template's own title color)
NAVY        = RGBColor(0x14, 0x1C, 0x78)   # theme accent1
BLEED       = RGBColor(0xB2, 0x4F, 0x2C)
BLEED_SOFT  = RGBColor(0xF1, 0xDC, 0xD1)
ISCH        = RGBColor(0x21, 0x5F, 0x73)
ISCH_SOFT   = RGBColor(0xD8, 0xE5, 0xE8)

SERIF = "Times New Roman"   # theme major font
SANS  = "Arial"              # theme minor font

prs = Presentation(TEMPLATE)
SW, SH = 13.333, 7.5
master = prs.slide_masters[0]
LAYOUT_TITLE   = master.slide_layouts[0]   # "1_Diapositiva titolo"
LAYOUT_CONTENT = master.slide_layouts[4]   # "Solo titolo"
LAYOUT_BLANK   = master.slide_layouts[5]   # "Vuoto"

TOTAL_SLIDES = 29
RUNNING_TITLE = "CATE Estimation and Use in the MASTER DAPT Study · T. Pioda"

# ------------------------------------------------------------- primitives --
_slide_no = [1]  # slide 1 is the template's own title slide, added directly below

def _copy_placeholder(slide, layout, idx):
    for ph in layout.placeholders:
        if ph.placeholder_format.idx == idx:
            slide.shapes._spTree.append(copy.deepcopy(ph._element))
            return

def new_slide(layout=LAYOUT_CONTENT, footer=RUNNING_TITLE):
    s = prs.slides.add_slide(layout)
    _slide_no[0] += 1
    for idx in (10, 11, 12):
        _copy_placeholder(s, layout, idx)
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 12:
            # Bake a static page number rather than relying on a live field,
            # so it renders correctly in every viewer, not only ones that
            # recompute PowerPoint fields on open.
            ph.text_frame.text = str(_slide_no[0])
    if footer:
        for ph in s.placeholders:
            if ph.placeholder_format.idx == 11:
                ph.text_frame.text = footer
                for p in ph.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)
                        r.font.name = SANS
                        r.font.color.rgb = INK_FAINT
    return s

def set_title(slide, text, size=26):
    ph = slide.placeholders[0]
    ph.text_frame.word_wrap = True
    p = ph.text_frame.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.name = SERIF; r.font.bold = True
    r.font.color.rgb = TITLE_BLUE
    return ph

def _runs_from_markup(p, text, base_size, base_color, bold_color=None, font=SANS):
    import re
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if not part:
            continue
        r = p.add_run()
        if part.startswith("**") and part.endswith("**"):
            r.text = part[2:-2]
            r.font.bold = True
            r.font.color.rgb = bold_color or INK
        else:
            r.text = part
            r.font.color.rgb = base_color
        r.font.size = Pt(base_size)
        r.font.name = font

def add_eyebrow(slide, label, color=NAVY, x=0.75, y=0.88):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(9), Inches(0.3))
    tf = tb.text_frame; tf.word_wrap = False
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = "●  "
    r1.font.size = Pt(10); r1.font.name = SANS; r1.font.color.rgb = color
    r2 = p.add_run(); r2.text = label.upper()
    r2.font.size = Pt(10); r2.font.name = SANS; r2.font.bold = True; r2.font.color.rgb = color
    return tb

def add_kicker(slide, text, x=0.75, y=1.55, w=9.5, h=0.6, size=15, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.3
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.italic = True; r.font.name = SERIF; r.font.color.rgb = INK_SOFT
    return tb

def add_bullets(slide, items, x, y, w, h, dot_color=INK_FAINT, size=14, gap=12, line_spacing=1.15):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        p.line_spacing = line_spacing
        rdot = p.add_run(); rdot.text = "●  "
        rdot.font.size = Pt(size - 3); rdot.font.name = SANS; rdot.font.color.rgb = dot_color
        _runs_from_markup(p, item, size, INK_SOFT, bold_color=INK)
    return tb

def add_card(slide, x, y, w, h, fill=SURFACE, border=LINE):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        shp.adjustments[0] = 0.045
    except Exception:
        pass
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = border; shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp

def add_mini_card(slide, x, y, w, h, header, body, header_color=NAVY, dashed=False):
    fill = SURFACE2 if dashed else SURFACE
    add_card(slide, x, y, w, h, fill=fill)
    tb = slide.shapes.add_textbox(Inches(x + 0.16), Inches(y + 0.13), Inches(w - 0.32), Inches(h - 0.26))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = header.upper()
    r.font.size = Pt(9.5); r.font.bold = True; r.font.name = SANS; r.font.color.rgb = header_color
    p2 = tf.add_paragraph(); p2.space_before = Pt(6); p2.line_spacing = 1.15
    _runs_from_markup(p2, body, 11, INK_SOFT, bold_color=INK)
    return tb

def add_callout(slide, text, x, y, w, h, color=INK, size=16):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.035), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background(); bar.shadow.inherit = False
    tb = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y - 0.06), Inches(w - 0.2), Inches(h + 0.15))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.line_spacing = 1.2
    _runs_from_markup(p, text, size, INK, bold_color=INK, font=SERIF)
    for r in p.runs:
        r.font.italic = True
    return tb

def add_stat(slide, x, y, w, num, label, color=INK):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.62))
    tf = tb.text_frame; tf.word_wrap = False
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = num
    r.font.size = Pt(27); r.font.name = SERIF; r.font.bold = True; r.font.color.rgb = color
    tb2 = slide.shapes.add_textbox(Inches(x), Inches(y + 0.56), Inches(w), Inches(0.4))
    tf2 = tb2.text_frame; tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run(); r2.text = label.upper()
    r2.font.size = Pt(9); r2.font.name = SANS; r2.font.color.rgb = INK_FAINT
    return tb

def add_table_rows(slide, x, y, w, col_fracs, headers, rows, row_h=0.42, row_colors=None):
    xs = [x]
    for f in col_fracs:
        xs.append(xs[-1] + f * w)
    for i, htxt in enumerate(headers):
        tb = slide.shapes.add_textbox(Inches(xs[i]), Inches(y), Inches(xs[i+1]-xs[i]), Inches(0.3))
        tf = tb.text_frame; tf.word_wrap = False
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = htxt.upper()
        r.font.size = Pt(9.5); r.font.bold = True; r.font.name = SANS; r.font.color.rgb = INK_FAINT
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.34), Inches(w), Pt(1))
    ln.fill.solid(); ln.fill.fore_color.rgb = LINE; ln.line.fill.background(); ln.shadow.inherit = False
    cy = y + 0.42
    for ri, row in enumerate(rows):
        color = INK
        if row_colors and ri < len(row_colors) and row_colors[ri]:
            color = row_colors[ri]
        for ci, val in enumerate(row):
            tb = slide.shapes.add_textbox(Inches(xs[ci]), Inches(cy), Inches(xs[ci+1]-xs[ci]), Inches(row_h - 0.06))
            tf = tb.text_frame; tf.word_wrap = False
            p = tf.paragraphs[0]
            r = p.add_run(); r.text = val
            r.font.size = Pt(12.5); r.font.name = SANS
            r.font.color.rgb = color
            r.font.bold = (color != INK)
        ln2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(cy + row_h - 0.06), Inches(w), Pt(0.75))
        ln2.fill.solid(); ln2.fill.fore_color.rgb = LINE; ln2.line.fill.background(); ln2.shadow.inherit = False
        cy += row_h
    return cy

def add_image_card(slide, path, x, y, maxw, maxh, caption=None, center_in_w=None):
    im = PILImage.open(path)
    iw, ih = im.size
    ar = iw / ih
    w = maxw
    h = w / ar
    if h > maxh:
        h = maxh
        w = h * ar
    pad = 0.16
    card_w = w + pad * 2
    card_h = h + pad * 2
    box_w = center_in_w if center_in_w else maxw
    cx = x + (box_w - card_w) / 2
    add_card(slide, cx, y, card_w, card_h, fill=SURFACE)
    slide.shapes.add_picture(path, Inches(cx + pad), Inches(y + pad), width=Inches(w), height=Inches(h))
    bottom = y + card_h
    if caption:
        cap_h = min(0.55, max(0.3, 7.35 - (bottom + 0.06)))
        cap = slide.shapes.add_textbox(Inches(cx), Inches(bottom + 0.06), Inches(card_w), Inches(cap_h))
        tf = cap.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.line_spacing = 1.15
        r = p.add_run(); r.text = caption
        r.font.size = Pt(8.5); r.font.name = SANS; r.font.color.rgb = INK_FAINT
    return bottom

# =============================================================== SLIDE 1 ==
# Reuse the template's own pre-authored title slide instead of adding a new one.
s = prs.slides[0]
ph_title = s.placeholders[0]
ph_title.text_frame.paragraphs[0].add_run().text = "CATE Estimation and Use in the MASTER DAPT Study"
for r in ph_title.text_frame.paragraphs[0].runs:
    r.font.name = SERIF; r.font.bold = True; r.font.color.rgb = TITLE_BLUE

ph_sub = s.placeholders[1]
p = ph_sub.text_frame.paragraphs[0]
r = p.add_run()
r.text = ("Can we find, validate, and act on individual treatment effect heterogeneity in a "
          "randomized cardiology trial — and if we can't, why not?")
r.font.size = Pt(17); r.font.italic = True; r.font.name = SERIF; r.font.color.rgb = INK_SOFT

ph_body = s.placeholders[13]
ph_body.text_frame.clear()
meta = [("Candidate", "Tommaso Pioda"), ("Supervisor", "Prof. Laura Azzimonti"),
        ("Co-supervisor", "Gabriele Maroni"), ("Client", "EOC"), ("Programme", "Data Science & AI")]
tf = ph_body.text_frame
for i, (k, v) in enumerate(meta):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    r1 = p.add_run(); r1.text = f"{k.upper()}:  "
    r1.font.size = Pt(11); r1.font.name = SANS; r1.font.bold = True; r1.font.color.rgb = INK_FAINT
    r2 = p.add_run(); r2.text = v
    r2.font.size = Pt(12.5); r2.font.name = SANS; r2.font.color.rgb = INK

for ph in s.placeholders:
    if ph.placeholder_format.idx == 15:
        ph.text_frame.text = "SUPSI · DTI — Diploma Thesis"
        for r in ph.text_frame.paragraphs[0].runs:
            r.font.size = Pt(9); r.font.name = SANS; r.font.color.rgb = INK_FAINT

# =============================================================== SLIDE 2 ==
s = new_slide()
add_eyebrow(s, "Context")
set_title(s, "From the average effect to the individual decision")
add_bullets(s, [
    "Randomized controlled trials are the gold standard for causal claims — but large, "
    "costly, and slow, and they report a single number: the **average treatment effect**.",
    "An average effect conceals a clinically critical dimension: **not every patient "
    "benefits equally** from the same therapy.",
    "Precision medicine asks for more: the **Conditional Average Treatment Effect**, τ(x) "
    "— the effect expected for a patient with covariates x.",
], x=0.75, y=2.3, w=10.9, h=3.4, size=16, gap=20)
add_callout(s, "This thesis asks whether that promise — sharper decisions from the same trial "
               "data — holds for a real RCT: MASTER DAPT.", x=0.75, y=5.35, w=10.9, h=1.0)

# =============================================================== SLIDE 3 ==
s = new_slide()
add_eyebrow(s, "The problem")
set_title(s, "One trial, two ways to be wrong")
tb = s.shapes.add_textbox(Inches(0.75), Inches(2.2), Inches(6.6), Inches(1.6))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.line_spacing = 1.3
r = p.add_run()
r.text = ("MASTER DAPT randomized 4,579 high-bleeding-risk patients between abbreviated and "
          "prolonged DAPT, and concluded in favour of abbreviation — non-inferior on ischaemia, "
          "superior on bleeding.")
r.font.size = Pt(14); r.font.name = SANS; r.font.color.rgb = INK_SOFT
add_callout(s, "That is an average result. A favourable trial-level effect is compatible with "
               "subgroups that gain nothing — or are harmed.", x=0.75, y=4.0, w=6.6, h=1.4)
add_stat(s, 8.0, 2.35, 4.3, "↓ bleeding", "abbreviated DAPT", color=BLEED)
add_stat(s, 8.0, 3.3, 4.3, "↑ ischaemia risk", "if therapy is too short", color=ISCH)

# =============================================================== SLIDE 4 ==
s = new_slide()
add_eyebrow(s, "The problem")
set_title(s, "Why “treat the highest-risk patient” is not an answer")
add_bullets(s, [
    "A risk model ranks patients by **P(event | treatment received)** — one world, the one we observed.",
    "A treatment decision needs the **difference between two worlds**, and only one of them is ever observed for any patient. This is the fundamental problem of causal inference.",
    "The highest-risk patient may gain the most from shortening therapy, the least, or be harmed by it — risk alone cannot tell the three apart.",
], x=0.75, y=2.25, w=10.8, h=2.6, size=15, gap=16)
add_callout(s, "Risk does not rank benefit.", x=0.75, y=5.25, w=9, h=0.6, size=20)

# =============================================================== SLIDE 5 ==
s = new_slide()
add_eyebrow(s, "Research questions")
set_title(s, "Two questions, in order")
add_mini_card(s, 0.75, 2.35, 5.6, 2.0, "Q1 — Estimate & validate",
              "Can the individual treatment effect (CATE) be estimated at all — and, since the true "
              "individual effect is never observed, can it be **validated**?")
add_mini_card(s, 6.55, 2.35, 5.6, 2.0, "Q2 — Exploit",
              "If a credible effect exists, can enrolment be **concentrated** on the patients where "
              "it is visible — reaching the same precision with fewer patients?")
add_callout(s, "Five estimator families answer Q1. Five enrolment mechanisms answer Q2 — once Q1's "
               "answer changes what Q2 has to test.", x=0.75, y=4.8, w=10.8, h=1.0)

# =============================================================== SLIDE 6 ==
s = new_slide()
add_eyebrow(s, "Roadmap")
set_title(s, "How the talk gets from Q1 to Q2")
add_bullets(s, [
    "**Case study** — the MASTER DAPT cohort, and a predictive anomaly in its risk models.",
    "**Method** — five estimator families, validated against one standard: the noise floor.",
    "**Results** — the positive control, the negative result, the trade-off plane, decision "
    "value, and the smallest effect the data could still be hiding.",
    "**Why** — three converging explanations for the negative result.",
    "**Testing it end-to-end** — guided enrolment, an exploitability test, with Mechanism 1's "
    "own results.",
    "**Conclusions** — what holds, what doesn't, and where to ask the same question next.",
], x=0.75, y=2.3, w=10.9, h=4.4, size=15, gap=16)

# =============================================================== SLIDE 7 ==
s = new_slide()
add_eyebrow(s, "Case study — data")
set_title(s, "MASTER DAPT: the cohort")
bottom = add_table_rows(s, 0.75, 2.35, 5.6,
                         [0.55, 0.22, 0.23],
                         ["Endpoint (335 d)", "Events", "Rate"],
                         [["Bleeding", "506", "11.1%"],
                          ["BARC 2/3/5", "359", "7.8%"],
                          ["Myocardial infarction", "109", "2.4%"],
                          ["CV death", "81", "1.8%"],
                          ["Stroke", "35", "0.8%"]],
                         row_colors=[BLEED, None, ISCH, ISCH, ISCH])
add_callout(s, "4,579 randomized · 63 baseline covariates · T=1 = prolonged, so CATE always reads "
               "as the benefit of shortening.", x=0.75, y=bottom + 0.35, w=5.6, h=1.3)
add_image_card(s, IMG + "masterdapt_flow.png", 6.7, 1.9, 5.3, 4.15,
                caption="All 4,579 patients were enrolled because they were at high bleeding risk "
                        "— a design criterion, not a coincidence.")

# =============================================================== SLIDE 8 ==
s = new_slide()
add_eyebrow(s, "Case study — risk models")
set_title(s, "A predictive anomaly, left open on purpose")
add_image_card(s, IMG + "ROC_PR_out_of_fold.png", 0.75, 2.0, 7.0, 3.6,
                caption="Five endpoint-specific risk models, evaluated out-of-fold, threshold-free, "
                        "over the whole cohort.")
tb = s.shapes.add_textbox(Inches(8.15), Inches(2.05), Inches(4.4), Inches(1.0))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.line_spacing = 1.3
r = p.add_run(); r.text = "If risk tracked event frequency, bleeding — 506 events — should be the easiest endpoint to predict. It is the hardest."
r.font.size = Pt(12.5); r.font.name = SANS; r.font.color.rgb = INK_SOFT
add_stat(s, 8.15, 3.15, 1.8, "0.76", "CV death · 81 ev.", color=ISCH)
add_stat(s, 9.95, 3.15, 1.8, "0.72", "MI · 109 ev.", color=ISCH)
add_stat(s, 8.15, 4.2, 1.8, "0.62", "Bleeding · 506 ev.", color=BLEED)
add_callout(s, "Explained in full only in chapter 8 — the same cause reappears there for the causal result.",
            x=8.15, y=5.2, w=4.4, h=1.1)

# =============================================================== SLIDE 9 ==
s = new_slide()
add_eyebrow(s, "Method — estimation")
set_title(s, "Five estimator families, chosen to err differently")
cards5 = [
    ("Meta-learner", "Two outcome models, one per arm; CATE is their difference. Simple, inherits every model error."),
    ("Causal Forest", "Estimates heterogeneity directly; double ML removes confounding, honest splits avoid overstating structure."),
    ("Bayesian Causal Forest", "Regularizes the effect surface separately from the prognostic one — a confounder can't disguise itself as heterogeneity."),
    ("In-context model", "Pre-trained on synthetic causal problems, not tuned to this dataset — a bias-driven, variance-independent control."),
    ("Interaction Forest", "One model on [X, T, X·T]; pools information across arms, compresses weak interactions toward zero."),
    ("Why five?", "Not “which is best” — a result that survives all five stops being attributable to any single algorithm's quirks."),
]
gx, gy, gw, gh, gap = 0.75, 2.3, 3.63, 1.55, 0.22
for i, (h, b) in enumerate(cards5):
    col = i % 3; row = i // 3
    x = gx + col * (gw + gap)
    y = gy + row * (gh + gap)
    add_mini_card(s, x, y, gw, gh, h, b, dashed=(h == "Why five?"))

# ============================================================== SLIDE 10 ==
s = new_slide()
add_eyebrow(s, "Method — validation")
set_title(s, "Validating what can never be observed")
add_bullets(s, [
    "τ(x) is never observed for **any** patient — CATE can't be checked prediction-vs-truth like a risk model.",
    "Instead, validate the **ranking**: the TOC curve — treat the top q%, gain vs. treating everyone. **AUTOC** weights the head of the ranking, where the decision happens.",
    "Evaluation is doubly-robust, plus a group-calibration check between estimated and observed CATE bands.",
    "Hyperparameters are tuned on the **same** metric used to validate — the model is trained for exactly the objective that matters.",
], x=0.75, y=2.3, w=6.7, h=4.2, size=13.5, gap=14)
add_mini_card(s, 7.75, 2.65, 4.6, 3.1, "The noise floor",
              "Even a random CATE produces some ranking dispersion. So: permute the treatment labels, "
              "re-run the whole pipeline — that distribution is chance. A result counts only if it "
              "clears **this floor**, not zero.")

# ============================================================== SLIDE 11 ==
s = new_slide()
add_eyebrow(s, "Results — 1 of 2", color=BLEED)
set_title(s, "The positive control: the tools work")
tb = s.shapes.add_textbox(Inches(0.75), Inches(2.2), Inches(10.8), Inches(0.5))
tf = tb.text_frame; p = tf.paragraphs[0]
r = p.add_run(); r.text = "All five estimator families independently recover the trial's average effect."
r.font.size = Pt(15); r.font.name = SANS; r.font.color.rgb = INK_SOFT
add_stat(s, 0.75, 2.95, 4, "+0.045 → +0.055", "bleeding CATE, mean (abbreviating helps)", color=BLEED)
add_stat(s, 4.9, 2.95, 3, "≈ 0", "ischaemic CATE, mean", color=ISCH)
add_stat(s, 8.0, 2.95, 3, "5 / 5", "estimator families agree", color=INK)
add_callout(s, "Consistent with the published trial, across five families that make different "
               "mistakes. The average effect is not in question.", x=0.75, y=4.55, w=9.5, h=1.0, color=BLEED)

# ============================================================== SLIDE 12 ==
s = new_slide()
add_eyebrow(s, "Results — 2 of 2", color=ISCH)
set_title(s, "The negative result: no heterogeneity above chance")
add_bullets(s, [
    "Individual CATE dispersion sits **on** the permutation noise floor, not above it.",
    "AUTOC is compatible with zero; credible bands cover the ATE across the whole ranking.",
    "Holds for **all five** independent estimator families.",
], x=0.75, y=2.25, w=6.2, h=2.2, dot_color=ISCH, size=14, gap=16)
add_callout(s, "Not “I found nothing” — “I searched with tools proven to work, and there is nothing.”",
            x=0.75, y=4.6, w=6.2, h=1.4, color=ISCH)
add_image_card(s, IMG + "CATE_placebo_floor.png", 7.2, 2.0, 5.2, 4.5,
                caption="Observed sd(τ(x)) (red) overlapping the placebo floor (grey) — bleeding and BARC 2/3/5.")

# ============================================================== SLIDE 13 ==
s = new_slide()
add_eyebrow(s, "Method — combining two outcomes")
set_title(s, "Four zones, fixed clinical meanings")
add_bullets(s, [
    "Both axes are the effect of **abbreviating** DAPT, oriented so positive always means abbreviating helps.",
    "**Win-win** (top-right) and **lose-lose** (bottom-left): the two outcomes agree, so one trial-wide "
    "rule — abbreviate, or don't — is already correct for every patient in that zone.",
    "The two off-diagonal zones disagree: abbreviating helps one outcome and costs the other. That is a "
    "genuine **trade-off**, and its mirror image.",
    "Only there does the right choice depend on **who the patient is** — the zone a personalized rule "
    "has to find, tested against real data next.",
], x=0.75, y=2.2, w=5.85, h=3.35, size=13.5, gap=14)
add_callout(s, "The relative weight given to each outcome is a clinical judgement, not a statistical one — "
               "this thesis uses 0.2 death / 0.4 MI / 0.4 stroke, set by the study's clinicians.",
            x=0.75, y=5.75, w=5.85, h=1.2, size=13)
add_image_card(s, IMG + "CATE_tradeoff_zones.png", 7.0, 1.95, 5.3, 4.4,
                caption="Bottom-right = abbreviating prevents bleeds but costs ischaemic events — "
                        "“depends on the patient.”")

# ============================================================== SLIDE 14 ==
s = new_slide()
add_eyebrow(s, "Results — the trade-off plane")
set_title(s, "Where is the dilemma patient?")
add_image_card(s, IMG + "CATE_tradeoff_plane.png", 0.75, 2.0, 7.3, 4.6,
                caption="Every patient placed by bleeding-benefit (x) and ischaemic-benefit (y) of abbreviating.")
tb = s.shapes.add_textbox(Inches(8.4), Inches(2.15), Inches(4.2), Inches(1.4))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.line_spacing = 1.3
r = p.add_run(); r.text = "A personalized rule only pays off if a quadrant of real dilemma cases exists — large gain on one axis, real harm on the other."
r.font.size = Pt(13); r.font.name = SANS; r.font.color.rgb = INK_SOFT
add_callout(s, "Instead: a compact cloud around the trial's average effect. No quadrant to split on.",
            x=8.4, y=3.75, w=4.2, h=1.4, color=BLEED)

# ============================================================== SLIDE 15 ==
s = new_slide()
add_eyebrow(s, "Results — decision value")
set_title(s, "Even the weaker question gives the same answer")
tb = s.shapes.add_textbox(Inches(0.75), Inches(2.15), Inches(10.8), Inches(0.7))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
_runs_from_markup(p, "Not “does the model rank well” — **does following the ranking produce a measurable "
                      "clinical gain**, evaluated with cross-fitted AIPW policy value.", 14.5, INK_SOFT, bold_color=INK)
cards12 = [
    ("Best rule vs. everyone", "No rule beats “abbreviate for everyone.” Net-neutral."),
    ("7 pre-specified subgroups", "Bleeding risk, ACS, diabetes, renal function, age, anaemia, OAC — none exclude zero."),
    ("Trade-off frontier", "Risk-magnification frontier is empty at every bleeding/ischaemic weight."),
]
gw, gh, gap = 3.63, 2.0, 0.22
for i, (h, b) in enumerate(cards12):
    x = 0.75 + i * (gw + gap)
    add_mini_card(s, x, 3.15, gw, gh, h, b)

# ============================================================== SLIDE 16 ==
s = new_slide()
add_eyebrow(s, "Results — how sure are we?")
set_title(s, "The minimum detectable δ")
add_image_card(s, IMG + "delta_sensitivity.png", 0.75, 2.0, 6.85, 4.25,
                caption="Synthetic heterogeneity of known size δ, injected and re-estimated; "
                        "compared against the floor recomputed at the same δ.")
bottom = add_table_rows(s, 8.05, 2.15, 4.5,
                         [0.55, 0.45],
                         ["Endpoint", "δ*"],
                         [["Bleeding", "0.05 risk-diff"],
                          ["BARC 2/3/5", "0.05 risk-diff"],
                          ["MI", "OR 3.0"],
                          ["CV death", "OR 5.0"]],
                         row_colors=[BLEED, BLEED, ISCH, ISCH])
add_callout(s, "On bleeding, δ* ≈ the entire average effect. It was not found — that is not the same claim as it does not exist.",
            x=8.05, y=bottom + 0.35, w=4.5, h=1.7)

# ============================================================== SLIDE 17 ==
s = new_slide()
add_eyebrow(s, "Why — three converging explanations")
set_title(s, "Not just “nothing found” — a mechanism")
cards14 = [
    ("01 · Cancellation", "The aggregate ischaemic score sums MI and stroke, which move in **opposite directions** — the sum hides both."),
    ("02 · Range restriction", "Enrolling only high-bleeding-risk patients compresses bleeding-risk variance before any model sees a datum — the same cause behind the ch.5 anomaly."),
    ("03 · Rare events", "MI, death, stroke: 35–109 events. Three endpoints of five are **not measurable** here, not “flat.”"),
]
gw, gh, gap = 3.63, 2.1, 0.22
for i, (h, b) in enumerate(cards14):
    x = 0.75 + i * (gw + gap)
    add_mini_card(s, x, 2.3, gw, gh, h, b)
add_callout(s, "The same feature of the trial's design — range restriction on bleeding risk — "
               "explains both the ch.5 predictive anomaly and the ch.6 absence of heterogeneity.",
            x=0.75, y=4.8, w=10.8, h=1.1)

# ============================================================== SLIDE 18 ==
s = new_slide()
add_eyebrow(s, "Testing it end-to-end")
set_title(s, "Guided enrolment: an exploitability test")
add_bullets(s, [
    "5 enrolment mechanisms — whole-pool, partially-random, duel, UCB1, Thompson — scored by "
    "several policies over **100 paired cohorts**.",
    "Mechanisms **do** move the enrolled cohort on the estimated-CATE plane, stably and "
    "distinguishably from random.",
    "On **observed outcomes**, the shift is ≈ zero — even for the adaptive bandits.",
], x=0.75, y=2.2, w=6.7, h=2.9, size=14, gap=16)
add_mini_card(s, 7.75, 2.45, 4.6, 2.5, "ρ ≈ 0.91",
              "Win-win score correlates with the bleeding axis alone — confirms the cancellation.")
tb = s.shapes.add_textbox(Inches(7.9), Inches(5.05), Inches(4.4), Inches(0.5))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run(); r.text = "The adaptive cohort still loses to a plain randomized trial at equal enrolled n."
r.font.size = Pt(11.5); r.font.name = SANS; r.font.color.rgb = INK_SOFT
add_callout(s, "The machinery selects well on a signal that does not exist.",
            x=0.75, y=5.65, w=10.8, h=0.7, size=18)

# ============================================================== SLIDE 19 ==
s = new_slide()
add_eyebrow(s, "Testing it end-to-end — mechanism 1")
set_title(s, "Mechanism 1: it moves the cohort, on both axes")
tb = s.shapes.add_textbox(Inches(0.75), Inches(2.15), Inches(10.9), Inches(0.5))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run()
r.text = ("Sees the whole remaining pool every round; each policy vs. the random control, "
          "at n = 3,500 (100 seeded runs).")
r.font.size = Pt(13); r.font.name = SANS; r.font.italic = True; r.font.color.rgb = INK_FAINT
bottom = add_table_rows(s, 0.75, 2.85, 10.9,
                         [0.30, 0.175, 0.175, 0.175, 0.175],
                         ["Policy", "Isch. Δ (pp)", "p", "Bleed Δ (pp)", "p"],
                         [["Conflict",          "−0.16", "<0.001", "−0.07", "0.139"],
                          ["Net-benefit",       "−0.01", "0.944",  "+0.19", "<0.001"],
                          ["Angle net-benefit", "−0.02", "0.770",  "+0.17", "0.004"],
                          ["Angle conflict",    "−0.20", "<0.001", "−0.13", "0.003"],
                          ["−isch",             "−0.12", "0.013",  "−0.14", "0.006"],
                          ["+isch",             "+0.03", "0.245",  "+0.17", "0.002"]],
                         row_h=0.38)
add_callout(s, "Real, significant movement on both axes for most policies — but Mechanism 1's own "
               "model only modestly correlates with the tuned causal forest, weaker still on "
               "bleeding.", x=0.75, y=bottom + 0.3, w=10.9, h=1.0)

# ============================================================== SLIDE 20 ==
s = new_slide()
add_eyebrow(s, "Testing it end-to-end — mechanism 1")
set_title(s, "The table's shape: crossing counts, separating rates")
add_image_card(s, IMG + "policy_comparison_ischaemic_patients_only_full.png", 0.75, 1.8, 11.0, 3.3)
add_bullets(s, [
    "**Top row**: cumulative counts — enrolled (red) rises, left-out (orange) falls, simply "
    "because the enrolled arm keeps growing.",
    "**Bottom row** is what matters: ischaemic **event rate (%)**, enrolled vs. left-out, "
    "95% CI over 100 runs.",
], x=0.75, y=5.35, w=5.6, h=1.75, size=12, gap=6)
add_bullets(s, [
    "Conflict and Angle conflict separate visibly and stay separated — matches the table's "
    "p < 0.001.",
    "Net-benefit, Angle net-benefit and random overlap fully — no separation, Δ ≈ 0.",
], x=6.55, y=5.35, w=5.6, h=1.75, size=12, gap=6)

# ============================================================== SLIDE 21 ==
s = new_slide()
add_eyebrow(s, "Testing it end-to-end — scoring policies")
set_title(s, "What each policy actually selects, on the same plane")
add_image_card(s, IMG + "selection_frequency_by_policy_full.png", 0.75, 1.9, 6.7, 4.6,
                caption="500 refits of the same forest — brighter = a patient is picked more "
                        "often by that policy's score.")
add_bullets(s, [
    "Same bleed/isch trade-off plane as slides 13-14, one panel per scoring policy, colored "
    "by how often 500 independent refits pick that patient.",
    "Each policy's own decision rule shows up as a sharp boundary — Conflict and Angle "
    "conflict cut near the **−45°** diagonal, Net-benefit and Angle net-benefit near "
    "**+45°**, ±isch cut along the y-axis alone.",
    "The boundary is stable, not noisy — confirms slide 18's claim that mechanisms move the "
    "cohort **stably and distinguishably from random**.",
    "Stable is not the same as useful: slide 14 already showed this plane has no dilemma "
    "quadrant for any of these boundaries to isolate.",
], x=7.75, y=2.0, w=4.85, h=4.3, size=12.5, gap=14)

# ============================================================== SLIDE 22 ==
s = new_slide()
add_eyebrow(s, "Testing it end-to-end — mechanism 1")
set_title(s, "Same test, scored on MI alone")
add_image_card(s, IMG + "policy_comparison_ischaemic_patients_only_mi.png", 0.75, 1.8, 11.0, 3.3)
add_bullets(s, [
    "Same Mechanism 1 test as the last two slides, but the score now sees **only** the "
    "MI component (weight 1 on MI, 0 on death/stroke) instead of the 0.2/0.4/0.4 "
    "composite.",
    "The outcome plotted is unchanged — the composite ischaemic rate (death + MI + stroke) "
    "— so this checks whether MI-only scoring still moves the same real endpoint.",
], x=0.75, y=5.35, w=5.6, h=1.75, size=12, gap=6)
add_bullets(s, [
    "Net-benefit and Angle net-benefit now separate too — flat under the composite score "
    "on slide 20, but enrolled and left-out visibly diverge here.",
    "+isch shows the largest gap of any policy or weighting tried — exploratory, same "
    "100-run averaging as the main result, but not yet reduced to a significance table "
    "like slide 19's.",
], x=6.55, y=5.35, w=5.6, h=1.75, size=12, gap=6)

# ============================================================== SLIDE 23 ==
s = new_slide()
add_eyebrow(s, "Testing it end-to-end — mechanism 1")
set_title(s, "Same test, scored on stroke alone")
add_image_card(s, IMG + "policy_comparison_ischaemic_patients_only_stroke.png", 0.75, 1.8, 11.0, 3.3)
add_bullets(s, [
    "Same Mechanism 1 test as the last three slides, but the score now sees **only** the "
    "stroke component (weight 1 on stroke, 0 on death/MI) instead of the 0.2/0.4/0.4 "
    "composite.",
    "The outcome plotted is unchanged — the composite ischaemic rate (death + MI + stroke) "
    "— so this checks whether stroke-only scoring still moves the same real endpoint.",
], x=0.75, y=5.35, w=5.6, h=1.75, size=12, gap=6)
add_bullets(s, [
    "Net-benefit and Angle net-benefit separate here too, much like the MI-only weighting "
    "on the previous slide — flat only under the full composite score on slide 20.",
    "+isch again shows the largest gap — exploratory, same 100-run averaging as the main "
    "result, but not yet reduced to a significance table like slide 19's.",
], x=6.55, y=5.35, w=5.6, h=1.75, size=12, gap=6)

# ============================================================== SLIDE 24 ==
s = new_slide()
add_eyebrow(s, "Conclusions — what holds")
set_title(s, "Six claims, each independently supported")
cardsC = [
    ("Average effect", "Confirmed by five independent estimators."),
    ("No heterogeneity", "Validated against the **noise floor**, not against zero."),
    ("Confirmed twice", "ch.7's observed events reach the same conclusion, without touching a single estimated CATE."),
    ("Cohort, not tools", "The limit is shown twice — by the positive control and by δ*."),
    ("Known mechanism", "**HBR range restriction** — also explains the ch.5 predictive anomaly."),
    ("Two reusable lessons", "Selecting ≠ allocating; linear scores collapse onto the higher-variance axis."),
]
gx, gy, gw, gh, gap = 0.75, 2.3, 3.63, 1.55, 0.22
for i, (h, b) in enumerate(cardsC):
    col = i % 3; row = i // 3
    x = gx + col * (gw + gap)
    y = gy + row * (gh + gap)
    add_mini_card(s, x, y, gw, gh, h, b)

# ============================================================== SLIDE 25 ==
s = new_slide()
add_eyebrow(s, "Conclusions — limits")
set_title(s, "What doesn't hold, said before it's asked")
add_bullets(s, [
    "δ* = 0.05 is **large** — these data exclude heterogeneity the size of the whole ATE, not any heterogeneity.",
    "**3 of 5 endpoints** (35–109 events) are not assessable at all — “not measurable,” not “absent.”",
    "Stroke sensitivity analysis is not yet interpretable.",
    "Subgroup analyses are in-sample and descriptive; no external cohort; multiplicity declared but not formally corrected.",
], x=0.75, y=2.25, w=10.9, h=4.4, size=15, gap=18)

# ============================================================== SLIDE 26 ==
s = new_slide()
add_eyebrow(s, "Where to ask the same question next")
set_title(s, "The infrastructure is ready — the cohort wasn't")
add_bullets(s, [
    "A cohort **not** selected purely on bleeding risk — where the bleeding axis keeps its full variance.",
    "Longer follow-up, or a higher-ischaemic-risk population — 35 stroke events measure nothing, whatever the estimator.",
    "A trial **sized on the interaction**, not the average effect — δ* is exactly the tool to pre-register that target.",
], x=0.75, y=2.25, w=10.9, h=2.6, size=15, gap=18)
add_callout(s, "The guided-enrolment machinery of ch.7 is built, measured, and already equipped with "
               "its two known failure modes — ready for a cohort where heterogeneity actually exists.",
            x=0.75, y=5.05, w=10.9, h=1.3)

# ============================================================== SLIDE 27 ==
s = new_slide(layout=LAYOUT_BLANK, footer=None)
tb = s.shapes.add_textbox(Inches(0), Inches(2.7), Inches(SW), Inches(0.5))
tf = tb.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r1 = p.add_run(); r1.text = "●  "
r1.font.size = Pt(11); r1.font.color.rgb = INK_FAINT
r2 = p.add_run(); r2.text = "THANK YOU"
r2.font.size = Pt(11); r2.font.bold = True; r2.font.name = SANS; r2.font.color.rgb = INK_FAINT
tb = s.shapes.add_textbox(Inches(0), Inches(3.15), Inches(SW), Inches(1.2))
tf = tb.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Questions"
r.font.size = Pt(46); r.font.name = SERIF; r.font.bold = True; r.font.color.rgb = TITLE_BLUE
add_kicker(s, "Tommaso Pioda · CATE Estimation and Use in the MASTER DAPT Study\n"
              "Supervisor: Prof. Laura Azzimonti — Co-supervisor: Gabriele Maroni",
           x=0, y=4.35, w=SW, h=1.0, size=14, align=PP_ALIGN.CENTER)

# ============================================================== SLIDE 28 ==
s = new_slide()
add_eyebrow(s, "Backup — glossary")
set_title(s, "Terms used without stopping to define them")
cardsG = [
    ("AUTOC", "Area under the TOC curve, weighted toward its top — the region a treat-the-top-q% decision actually uses."),
    ("TOC curve", "Targeting Operator Characteristic: gain from treating the top q% ranked by estimated CATE, vs. treating everyone."),
    ("Doubly-robust", "Consistent if either the outcome model or the propensity model is correct, not both — protects the evaluation from either being wrong."),
    ("AIPW policy value", "Augmented inverse-propensity-weighted estimate of the outcome under a decision rule — the metric behind ch.6's decision-value results."),
    ("Honest splits", "A causal forest splits the covariate space on one subsample and estimates effects on another, so the same data never both proposes and confirms a split."),
    ("Group calibration", "Checks that patients binned by estimated CATE show the same ordering in their observed outcomes — a coarse validation the AUTOC ranking can't provide alone."),
]
gx, gy, gw, gh, gap = 0.75, 2.3, 3.63, 1.55, 0.22
for i, (h, b) in enumerate(cardsG):
    col = i % 3; row = i // 3
    x = gx + col * (gw + gap)
    y = gy + row * (gh + gap)
    add_mini_card(s, x, y, gw, gh, h, b)

# ============================================================== SLIDE 29 ==
s = new_slide()
add_eyebrow(s, "Backup — additional robustness checks")
set_title(s, "Checks already run, not shown in the main talk")
cardsR = [
    ("CF flexibility presets", "Regularized, medium, and flexible causal-forest fits all collapse to zero heterogeneity (ch.6) — not an artefact of one hyperparameter choice."),
    ("Covariate discretization", "Coarsening every covariate into 4 ordinal bins and refitting reaches the same conclusion (ch.6) — heterogeneity isn't hiding in a continuous-vs-binned assumption."),
    ("Stroke, specifically", "35 events fail the sensitivity analysis's own null-calibration check (ch.8) — no interpretable δ* exists for it, not \"not yet run.\""),
]
gw, gh, gap = 3.63, 2.1, 0.22
for i, (h, b) in enumerate(cardsR):
    x = 0.75 + i * (gw + gap)
    add_mini_card(s, x, 2.3, gw, gh, h, b)

# ------------------------------------------------------------------ save --
prs.save(OUT)
print("Saved:", OUT)
print("Slides:", len(list(prs.slides)))
